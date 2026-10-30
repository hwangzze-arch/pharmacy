# -*- coding: utf-8 -*-
"""실험1 뷰렛의 검정 결과 엑셀 (Table I · 구간별 보정치 · 검정곡선 차트)"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.comments import Comment

OUT = sys.argv[1]
READ = [0.00, 10.00, 20.00, 30.00, 39.98, 50.00]
MASS = [59.0550, 69.0390, 79.0397, 89.0361, 98.9286, 108.8970]

F = "Arial"
thin = Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="EEF0F8")
inp = Font(name=F, color="0000FF")
blk = Font(name=F)
bold = Font(name=F, bold=True)
c_ = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "Table I"
ws["A1"] = "Table I. 50 mL 뷰렛의 검정 결과 (수온 19.0 °C, 1회 측정)"
ws["A1"].font = Font(name=F, bold=True, size=12)
ws["A2"] = "수온 (°C)"; ws["B2"] = 19.0
ws["A3"] = "물의 밀도 ρ (g/mL)"; ws["B3"] = 0.99842
ws["C3"] = "유인물 5쪽 Table 1.2 (19.0 °C)"
for a in ("A2", "A3"): ws[a].font = bold
ws["B2"].font = inp; ws["B3"].font = inp
ws["B2"].number_format = "0.0"; ws["B3"].number_format = "0.00000"
ws["C3"].font = Font(name=F, italic=True, color="666666")

H = ["뷰렛 눈금 (mL)", "플라스크 + 물 (g)", "물 무게 (g)", "실제 부피 (mL)", "보정치 (mL)"]
for j, h in enumerate(H):
    c = ws.cell(row=5, column=1 + j, value=h)
    c.font = bold; c.fill = hdr_fill; c.alignment = c_; c.border = box
for i, (r, m) in enumerate(zip(READ, MASS)):
    row = 6 + i
    ws.cell(row=row, column=1, value=r).number_format = "0.00"
    ws.cell(row=row, column=2, value=m).number_format = "0.0000"
    ws.cell(row=row, column=3, value=f"=B{row}-$B$6").number_format = "0.0000"
    ws.cell(row=row, column=4, value=f"=ROUND(C{row}/$B$3,2)").number_format = "0.00"
    ws.cell(row=row, column=5, value=f"=ROUND(D{row}-A{row},2)").number_format = '+0.00;-0.00;0.00'
    for j in range(5):
        c = ws.cell(row=row, column=1 + j)
        c.border = box; c.alignment = c_
        c.font = inp if j < 2 else blk
ws["A6"].comment = Comment("0.00 mL에서 시작 (빈 플라스크)", "Claude")
ws["B10"].comment = Comment("40 mL 측정 도중 소량 흘림 (실험자 기록)", "Claude")

ws["A13"] = "구간별 보정치 (Table II)"; ws["A13"].font = bold
H2 = ["구간 (mL)", "눈금 차 (mL)", "실제 부피 차 (mL)", "구간 보정치 (mL)"]
for j, h in enumerate(H2):
    c = ws.cell(row=14, column=1 + j, value=h)
    c.font = bold; c.fill = hdr_fill; c.alignment = c_; c.border = box
for i in range(5):
    row, a, b = 15 + i, 6 + i, 7 + i
    ws.cell(row=row, column=1, value=f'=TEXT(A{a},"0")&" → "&TEXT(A{b},"0")')
    ws.cell(row=row, column=2, value=f"=A{b}-A{a}").number_format = "0.00"
    ws.cell(row=row, column=3, value=f"=D{b}-D{a}").number_format = "0.00"
    ws.cell(row=row, column=4, value=f"=ROUND(C{row}-B{row},2)").number_format = '+0.00;-0.00;0.00'
    for j in range(4):
        c = ws.cell(row=row, column=1 + j); c.border = box; c.alignment = c_; c.font = blk

ws["A21"] = "판정"; ws["A21"].font = bold
ws["B21"] = '=IF(MAX(ABS(MAX(E6:E11)),ABS(MIN(E6:E11)))<=0.1,"모든 보정치가 ±0.1 mL 이내 → 허용오차 만족","±0.1 mL 초과 구간 있음")'
ws["A22"] = "보정치 범위"; ws["A22"].font = bold
ws["B22"] = '=TEXT(MIN(E6:E11),"+0.00;-0.00;0.00")&" ~ "&TEXT(MAX(E6:E11),"+0.00;-0.00;0.00")&" mL"'
ws["A24"] = "파란 글자 = 측정값(입력), 검은 글자 = 수식. 물 무게 = 총무게 − 빈 플라스크, 실제 부피 = 물 무게 ÷ ρ, 보정치 = 실제 부피 − 눈금."
ws["A24"].font = Font(name=F, italic=True, color="666666")
for a in ("B21", "B22", "A21", "A22"): ws[a].font = bold if a[0] == "A" else blk

for col, w in zip("ABCDE", (21, 18, 16, 17, 16)):
    ws.column_dimensions[col].width = w
ws.row_dimensions[5].height = 30

ch = ScatterChart()
ch.title = "Figure 1. 50 mL 뷰렛의 검정곡선 (19.0 °C)"
ch.style = 13
ch.x_axis.title = "뷰렛 눈금 (mL)"
ch.y_axis.title = "보정치 (mL)"
ch.x_axis.scaling.min = 0; ch.x_axis.scaling.max = 50; ch.x_axis.majorUnit = 10
ch.y_axis.scaling.min = -0.12; ch.y_axis.scaling.max = 0.12; ch.y_axis.majorUnit = 0.02
ch.y_axis.number_format = "+0.00;-0.00;0.00"
ch.x_axis.delete = False; ch.y_axis.delete = False
xs = Reference(ws, min_col=1, min_row=6, max_row=11)
ys = Reference(ws, min_col=5, min_row=5, max_row=11)
s = Series(ys, xs, title_from_data=True)
s.marker.symbol = "circle"; s.marker.size = 7
s.graphicalProperties.line.solidFill = "4651C8"; s.graphicalProperties.line.width = 22000
s.smooth = False
ch.series.append(s)
ws["A26"] = "그래프용 허용오차 선"; ws["A26"].font = bold
ws["A27"] = "x (mL)"; ws["B27"] = 0; ws["C27"] = 50
ws["A28"] = "+0.10"; ws["B28"] = 0.1; ws["C28"] = 0.1
ws["A29"] = "-0.10"; ws["B29"] = -0.1; ws["C29"] = -0.1
for r in (27, 28, 29):
    for col in "ABC": ws[f"{col}{r}"].font = Font(name=F, color="666666")
tx = Reference(ws, min_col=2, max_col=3, min_row=27)
for r in (28, 29):
    t = Series(Reference(ws, min_col=2, max_col=3, min_row=r), tx, title=f"허용오차 {ws[f'A{r}'].value}")
    t.marker.symbol = "none"; t.smooth = False
    t.graphicalProperties.line.solidFill = "1F7A3F"; t.graphicalProperties.line.dashStyle = "dash"; t.graphicalProperties.line.width = 15000
    ch.series.append(t)
ch.x_axis.tickLblPos = "low"; ch.y_axis.tickLblPos = "low"
ch.x_axis.number_format = "0"
ch.legend = None
ch.width = 16; ch.height = 9
ws.add_chart(ch, "G2")
wb.save(OUT)
